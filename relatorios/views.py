from django.shortcuts import render
from django.views.generic import TemplateView, ListView
from django.http import HttpResponse
from django.db.models import Count, Q, F, DecimalField, Case, When, Value, Sum
from django.db.models.functions import Coalesce
from datetime import date, timedelta
from accounts.mixins import CoordenadorRequiredMixin
from alunos.models import Aluno, Matricula
from escola.models import Turma
from frequencia.models import Frequencia
from monitoramento.services import FrequenciaService


class RelatorioAlunoView(CoordenadorRequiredMixin, TemplateView):
    """Individual student report"""
    template_name = 'relatorios/relatorio_aluno.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        aluno_id = self.request.GET.get('aluno_id')
        data_inicio = self.request.GET.get('data_inicio')
        data_fim = self.request.GET.get('data_fim')
        
        if aluno_id:
            try:
                aluno = Aluno.objects.get(id=aluno_id)
                context['aluno'] = aluno
                
                # Get frequency data
                frequencias = Frequencia.objects.filter(aluno=aluno)
                
                if data_inicio:
                    frequencias = frequencias.filter(data__gte=data_inicio)
                if data_fim:
                    frequencias = frequencias.filter(data__lte=data_fim)
                
                context['frequencias'] = frequencias.select_related('turma').order_by('-data')
                
                # Calculate statistics
                total = frequencias.count()
                presentes = frequencias.filter(status='P').count()
                ausentes = frequencias.filter(status='A').count()
                justificadas = frequencias.filter(status='J').count()
                
                context['total_aulas'] = total
                context['presentes'] = presentes
                context['ausentes'] = ausentes
                context['justificadas'] = justificadas
                
                if total > 0:
                    context['percentual_frequencia'] = ((presentes + justificadas) / total) * 100
                else:
                    context['percentual_frequencia'] = 0
                
            except Aluno.DoesNotExist:
                pass
        
        context['alunos'] = Aluno.objects.filter(status='ATIVO')
        context['aluno_selecionado'] = aluno_id
        context['data_inicio'] = data_inicio
        context['data_fim'] = data_fim
        
        return context


class RelatorioTurmaView(CoordenadorRequiredMixin, TemplateView):
    """Class attendance report"""
    template_name = 'relatorios/relatorio_turma.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        turma_id = self.request.GET.get('turma_id')
        data_inicio = self.request.GET.get('data_inicio')
        data_fim = self.request.GET.get('data_fim')
        
        if turma_id:
            try:
                turma = Turma.objects.get(id=turma_id)
                context['turma'] = turma
                
                # Get students in turma
                matriculas = Matricula.objects.filter(turma=turma, status='ATIVA')
                alunos = [m.aluno for m in matriculas]
                
                # Build report data
                relatorio_dados = []
                for aluno in alunos:
                    frequencias = Frequencia.objects.filter(aluno=aluno, turma=turma)
                    
                    if data_inicio:
                        frequencias = frequencias.filter(data__gte=data_inicio)
                    if data_fim:
                        frequencias = frequencias.filter(data__lte=data_fim)
                    
                    total = frequencias.count()
                    presentes = frequencias.filter(status='P').count()
                    ausentes = frequencias.filter(status='A').count()
                    justificadas = frequencias.filter(status='J').count()
                    
                    percentual = ((presentes + justificadas) / total * 100) if total > 0 else 0
                    
                    relatorio_dados.append({
                        'aluno': aluno,
                        'total': total,
                        'presentes': presentes,
                        'ausentes': ausentes,
                        'justificadas': justificadas,
                        'percentual': percentual
                    })
                
                context['relatorio_dados'] = sorted(relatorio_dados, key=lambda x: x['percentual'])
                
            except Turma.DoesNotExist:
                pass
        
        context['turmas'] = Turma.objects.filter(ativo=True)
        context['turma_selecionada'] = turma_id
        context['data_inicio'] = data_inicio
        context['data_fim'] = data_fim
        
        return context


class ExportarExcelView(CoordenadorRequiredMixin, TemplateView):
    """Export attendance data to Excel"""
    template_name = 'relatorios/exportar_excel.html'

    def post(self, request, *args, **kwargs):
        turma_id = request.POST.get('turma_id')
        data_inicio = request.POST.get('data_inicio')
        data_fim = request.POST.get('data_fim')
        
        if not turma_id:
            return self.get(request, *args, **kwargs)
        
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            
            turma = Turma.objects.get(id=turma_id)
            
            # Create workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Frequência"
            
            # Header
            ws['A1'] = "Relatório de Frequência"
            ws['A1'].font = Font(bold=True, size=14)
            ws['A2'] = f"Turma: {turma.nome} ({turma.serie.nome}) - {turma.ano}"
            ws['A3'] = f"Período: {data_inicio} a {data_fim}"
            
            # Column headers
            headers = ['Aluno', 'Total de Aulas', 'Presenças', 'Ausências', 'Justificadas', 'Frequência %']
            ws.append(headers)
            
            # Format headers
            header_fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            
            for cell in ws[5]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
            
            # Get data
            matriculas = Matricula.objects.filter(turma=turma, status='ATIVA')
            row = 6
            
            for matricula in matriculas:
                aluno = matricula.aluno
                frequencias = Frequencia.objects.filter(aluno=aluno, turma=turma)
                
                if data_inicio:
                    frequencias = frequencias.filter(data__gte=data_inicio)
                if data_fim:
                    frequencias = frequencias.filter(data__lte=data_fim)
                
                total = frequencias.count()
                presentes = frequencias.filter(status='P').count()
                ausentes = frequencias.filter(status='A').count()
                justificadas = frequencias.filter(status='J').count()
                percentual = ((presentes + justificadas) / total * 100) if total > 0 else 0
                
                ws[f'A{row}'] = aluno.nome
                ws[f'B{row}'] = total
                ws[f'C{row}'] = presentes
                ws[f'D{row}'] = ausentes
                ws[f'E{row}'] = justificadas
                ws[f'F{row}'] = f"{percentual:.2f}%"
                
                row += 1
            
            # Adjust column widths
            ws.column_dimensions['A'].width = 30
            ws.column_dimensions['B'].width = 15
            ws.column_dimensions['C'].width = 12
            ws.column_dimensions['D'].width = 12
            ws.column_dimensions['E'].width = 15
            ws.column_dimensions['F'].width = 15
            
            # Create response
            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="relatorio_frequencia_{turma.nome}_{date.today()}.xlsx"'
            wb.save(response)
            
            return response
            
        except Exception as e:
            return self.get(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['turmas'] = Turma.objects.filter(ativo=True)
        return context

